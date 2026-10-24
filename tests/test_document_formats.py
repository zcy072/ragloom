import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

try:
    import docx
    from langchain_core.documents import Document

    from routers import documents
    from services.documents import loaders as document_loader
    from services.documents.parser_modes import PDF_PARSER_MINERU_CLOUD
    from services.documents.types import SUPPORTED_DOCUMENT_EXTENSIONS
    from services.indexing import chunking
except ModuleNotFoundError:
    docx = None
    Document = None
    documents = None
    SUPPORTED_DOCUMENT_EXTENSIONS = ()
    document_loader = None
    chunking = None
    PDF_PARSER_MINERU_CLOUD = "mineru_cloud"


@unittest.skipUnless(
    docx and Document and documents and document_loader and chunking,
    "requires backend dependencies",
)
class DocumentFormatTest(unittest.TestCase):
    def test_document_option_contracts_have_stable_fields(self):
        parser_options = document_loader.list_document_parser_options()
        format_options = document_loader.list_document_format_options()

        self.assertEqual(
            {"id", "label", "available", "unavailable_reason", "extensions"},
            set(parser_options[0]),
        )
        mineru_option = next(
            option for option in parser_options if option["id"] == PDF_PARSER_MINERU_CLOUD
        )
        self.assertIn("unavailable_reason", mineru_option)
        self.assertTrue(all(set(option) == {"extension", "label"} for option in format_options))

    def test_markdown_routes_preserve_authored_sections_and_encoding(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "制度.md"
            source.write_text(
                "# 员工制度\n\n## 参考文献\n\n这是业务要求，不应删除。\n\n"
                "### 审批流程\n\n提交申请。\n\n| 步骤 | 要求 |\n| --- | --- |\n| 申请 | 提前三天 |",
                encoding="utf-8-sig",
            )

            routes = document_loader.load_document_routes(str(source), source.name)
            splits = chunking.split_structured_markdown_documents(
                routes.structured, chunk_size=800
            )

        self.assertEqual(len(routes.fixed), 1)
        self.assertNotIn("\ufeff", routes.fixed[0].page_content)
        combined = "\n".join(split.page_content for split in splits)
        self.assertIn("这是业务要求，不应删除", combined)
        approval = next(split for split in splits if "提交申请" in split.page_content)
        self.assertIn("制度 > 员工制度 > 参考文献", approval.metadata["contextual_title"])
        self.assertIn("表头：步骤 | 要求", combined)
        self.assertIn("申请 | 提前三天", combined)
        self.assertEqual(routes.fixed[0].metadata["parser_mode"], "native_text")
        self.assertEqual(routes.structured[0].metadata["parser_mode"], "native_text")

    def test_large_markdown_pipe_table_repeats_header_in_every_table_chunk(self):
        rows = "\n".join(
            f"| M{index:03d} | {index * 100} 元 | {'库存说明' * 20} | 华东仓 |"
            for index in range(24)
        )
        document = Document(
            page_content=(
                "## 产品库存\n\n"
                "表 1 产品库存\n\n"
                "| 型号 | 单价 | 库存 | 仓库 |\n"
                "| - | -: | :- | :-: |\n"
                f"{rows}"
            ),
            metadata={"source": "库存.md", "structured_profile": "native"},
        )

        splits = chunking.split_structured_markdown_documents(
            [document], chunk_size=800
        )
        table_splits = [split for split in splits if "表头：" in split.page_content]

        self.assertGreater(len(table_splits), 1)
        self.assertTrue(
            all("表头：型号 | 单价 | 库存 | 仓库" in split.page_content for split in table_splits)
        )
        combined = "\n".join(split.page_content for split in table_splits)
        for index in range(24):
            self.assertEqual(combined.count(f"M{index:03d} |"), 1)

    def test_markdown_pipe_table_normalization_skips_fenced_examples(self):
        document = Document(
            page_content=(
                "## 编写说明\n\n"
                "```markdown\n"
                "| 示例列 | 示例值 |\n"
                "| --- | --- |\n"
                "| A | B |\n"
                "```\n\n"
                "| 名称 | 表达式 |\n"
                "| --- | --- |\n"
                r"| 比率 | `完成数 \| 总数` |"
            ),
            metadata={"source": "说明.md", "structured_profile": "native"},
        )

        splits = chunking.split_structured_markdown_documents([document])
        combined = "\n".join(split.page_content for split in splits)

        self.assertIn("| 示例列 | 示例值 |", combined)
        self.assertNotIn("表头：示例列 | 示例值", combined)
        self.assertIn("表头：名称 | 表达式", combined)
        self.assertIn("比率 | `完成数 | 总数`", combined)

    def test_local_docx_routes_keep_headings_lists_and_table_shape(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "休假制度.docx"
            word = docx.Document()
            word.add_heading("员工休假制度", level=0)
            word.add_heading("申请流程", level=1)
            word.add_paragraph("提前提交申请", style="List Number")
            table = word.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "工龄"
            table.rows[0].cells[1].text = "年假"
            table.rows[1].cells[0].text = "满一年"
            table.rows[1].cells[1].text = "5天"
            word.save(source)

            routes = document_loader.load_document_routes(str(source), source.name)
            preview = document_loader.docx_preview_html(str(source))

        self.assertIn("工龄 | 年假", routes.fixed[0].page_content)
        structured = routes.structured[0].page_content
        self.assertIn("# 员工休假制度", structured)
        self.assertIn("## 申请流程", structured)
        self.assertIn("1. 提前提交申请", structured)
        self.assertIn("<table>", structured)
        self.assertIn("<table>", preview)
        self.assertIn("<h1>员工休假制度</h1>", preview)

    def test_local_docx_infers_numbered_headings_only_without_styled_headings(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "无样式制度.docx"
            word = docx.Document()
            chapter = word.add_paragraph()
            chapter.add_run("第一章 总则").bold = True
            word.add_paragraph("本章规定制度的适用范围。" * 40)
            section = word.add_paragraph()
            section.add_run("1.2 术语和定义").bold = True
            word.add_paragraph("本节定义业务处理中使用的术语。" * 40)
            word.add_paragraph("应当按照流程提交申请。", style="List Number")
            word.save(source)

            routes = document_loader.load_document_routes(str(source), source.name)
            splits = chunking.split_structured_markdown_documents(
                routes.structured, chunk_size=800
            )

        structured = routes.structured[0].page_content
        self.assertIn("## 第一章 总则", structured)
        self.assertIn("### 1.2 术语和定义", structured)
        self.assertIn("1. 应当按照流程提交申请。", structured)
        terms = next(split for split in splits if "本节定义业务处理" in split.page_content)
        self.assertIn("无样式制度 > 第一章 总则 > 1.2 术语和定义", terms.metadata["contextual_title"])

    def test_local_docx_does_not_infer_plain_numbered_sentences_as_headings(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "编号正文.docx"
            word = docx.Document()
            word.add_paragraph("1 申请人应当提交完整材料。")
            word.add_paragraph("2. 审批人应当在三个工作日内处理。")
            word.save(source)

            routes = document_loader.load_document_routes(str(source), source.name)

        self.assertNotIn("# 1 申请人", routes.structured[0].page_content)
        self.assertNotIn("# 2. 审批人", routes.structured[0].page_content)

    def test_docx_mineru_mode_uses_cloud_pages_and_saved_markdown(self):
        with patch.object(
            document_loader.mineru_client,
            "parse_document",
            return_value=[(0, "云端正文")],
        ) as parse_document, patch.object(
            document_loader.mineru_client,
            "read_saved_markdown",
            return_value="## 云端章节\n\n云端正文",
        ):
            routes = document_loader.load_document_routes(
                "/tmp/制度.docx", "制度.docx", PDF_PARSER_MINERU_CLOUD
            )

        parse_document.assert_called_once()
        self.assertEqual(routes.fixed[0].page_content, "云端正文")
        self.assertEqual(routes.structured[0].metadata["structured_profile"], "mineru")
        self.assertIn("云端章节", routes.structured[0].page_content)

    def test_public_format_contract_excludes_legacy_and_tabular_formats(self):
        self.assertEqual(
            set(SUPPORTED_DOCUMENT_EXTENSIONS),
            {".pdf", ".docx", ".md", ".txt", ".png", ".jpg", ".jpeg"},
        )
        self.assertTrue(
            {".doc", ".csv", ".xls", ".xlsx", ".xsd"}.isdisjoint(
                SUPPORTED_DOCUMENT_EXTENSIONS
            )
        )


if __name__ == "__main__":
    unittest.main()
