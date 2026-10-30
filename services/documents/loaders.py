"""按文件格式提取文档内容，并为 RAG 入库提供固定和结构化两条路线。"""

import html
import os
import re

from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from PIL import Image

from services.documents import mineru as mineru_client
from services.documents.contracts import DocumentFormatOption, DocumentParserOption
from services.documents.heading_patterns import (
    APPENDIX_CHILD_HEADING,
    APPENDIX_HEADING,
    CHAPTER_HEADING,
    CHINESE_MAJOR_HEADING,
    CHINESE_SUB_HEADING,
    STANDARD_HEADING,
    STANDARD_TOP_HEADING,
)
from services.documents.ocr import get_ocr_engine, ocr_image
from services.documents.parser_modes import (
    DEFAULT_PDF_PARSER_MODE,
    IMAGE_PARSER_LABELS,
    IMAGE_PARSER_VISION_MODEL,
    PDF_PARSER_LABELS,
    PDF_PARSER_MINERU_CLOUD,
    PDF_PARSER_NATIVE_TEXT,
    PDF_PARSER_PADDLE_OCR,
    normalize_document_parser_mode,
    normalize_pdf_parser_mode,
)
from services.documents.types import (
    IMAGE_EXTENSIONS,
    DocumentLoadError,
    DocumentRoutes,
)


def list_document_parser_options() -> list[DocumentParserOption]:
    """返回前端可选的解析器及可用状态，但绝不暴露 MinerU 密钥。"""
    mineru_ready = mineru_client.is_configured()
    return [
        {
            "id": PDF_PARSER_NATIVE_TEXT,
            "label": PDF_PARSER_LABELS[PDF_PARSER_NATIVE_TEXT],
            "available": True,
            "unavailable_reason": "",
            "extensions": [".pdf", ".docx"],
        },
        {
            "id": PDF_PARSER_MINERU_CLOUD,
            "label": PDF_PARSER_LABELS[PDF_PARSER_MINERU_CLOUD],
            "available": mineru_ready,
            "unavailable_reason": "未配置 MINERU_API_TOKEN" if not mineru_ready else "",
            "extensions": [".pdf", ".docx"],
        },
        {
            "id": PDF_PARSER_PADDLE_OCR,
            "label": PDF_PARSER_LABELS[PDF_PARSER_PADDLE_OCR],
            "available": True,
            "unavailable_reason": "",
            "extensions": [".pdf", ".png", ".jpg", ".jpeg"],
        },
        {
            "id": IMAGE_PARSER_VISION_MODEL,
            "label": IMAGE_PARSER_LABELS[IMAGE_PARSER_VISION_MODEL],
            "available": True,
            "unavailable_reason": "",
            "extensions": [".png", ".jpg", ".jpeg"],
        },
    ]


def list_document_format_options() -> list[DocumentFormatOption]:
    """返回前端支持上传的文件格式及其展示名称。"""
    return [
        {"extension": extension, "label": label}
        for extension, label in (
            (".pdf", "PDF"),
            (".docx", "Word"),
            (".md", "Markdown"),
            (".txt", "文本"),
            (".png", "PNG 图片"),
            (".jpg", "JPEG 图片"),
            (".jpeg", "JPEG 图片"),
        )
    ]


def _load_pdf_native_text(file_path: str, filename: str) -> list[Document]:
    """使用 PDF 原生文本层逐页提取内容，适合本身包含可复制文本的 PDF。"""
    try:
        loaded = PyPDFLoader(file_path).load()
    except Exception as exc:
        raise DocumentLoadError(f"原生文本提取失败：{exc}") from exc

    documents: list[Document] = []
    for page_number, document in enumerate(loaded):
        page = document.metadata.get("page", page_number)
        documents.append(
            Document(
                page_content=document.page_content,
                metadata={"source": filename, "page": page, "parser_mode": PDF_PARSER_NATIVE_TEXT},
            )
        )
    return documents


def _load_pdf_paddle_ocr(file_path: str, filename: str) -> list[Document]:
    """将 PDF 页面转图片并本地 OCR；识别为空时回退到本地文本层。"""
    ocr = get_ocr_engine()
    if not ocr:
        raise DocumentLoadError("PaddleOCR 不可用，请检查本地 OCR 依赖")
    try:
        from pdf2image import convert_from_path, pdfinfo_from_path

        fallback_pages = None
        documents: list[Document] = []
        total_pages = int(pdfinfo_from_path(file_path)["Pages"])
        for page_number in range(total_pages):
            # 一次只渲染一页，扫描件页数很多时内存不会随页数增长。
            images = convert_from_path(
                file_path, first_page=page_number + 1, last_page=page_number + 1
            )
            if not images:
                continue
            page_text = ocr_image(images[0], ocr)
            if not page_text.strip():
                # 该回退仍完全在本地执行，避免文字型 PDF 因 OCR 空结果而无法使用。
                fallback_pages = fallback_pages or PyPDFLoader(file_path).load()
                if page_number < len(fallback_pages):
                    page_text = fallback_pages[page_number].page_content
            documents.append(
                Document(
                    page_content=page_text,
                    metadata={"source": filename, "page": page_number, "parser_mode": PDF_PARSER_PADDLE_OCR},
                )
            )
        return documents
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"Paddle 本地 OCR 失败：{exc}") from exc


def _load_pdf_mineru_cloud(
    file_path: str,
    filename: str,
    *,
    reuse_saved_result: bool = False,
) -> list[Document]:
    """调用或复用 MinerU 云端结果，按页生成固定检索路线的文档。"""
    try:
        pages = mineru_client.read_saved_page_texts(file_path) if reuse_saved_result else []
        if not pages:
            pages = mineru_client.parse_document(
                file_path, result_dir=mineru_client.result_dir_for_document(file_path)
            )
    except mineru_client.MinerUError as exc:
        raise DocumentLoadError(str(exc)) from exc
    return [
        Document(
            page_content=page_text,
            metadata={"source": filename, "page": page_number, "parser_mode": PDF_PARSER_MINERU_CLOUD},
        )
        for page_number, page_text in pages
    ]


def load_mineru_markdown_document(
    file_path: str,
    filename: str,
    parser_mode: str | None,
) -> list[Document]:
    """将 MinerU 保存的完整 Markdown 作为独立的结构化索引输入。

    保留页码的 JSON 路线仍是主提取结果。该函数只读取本地缓存，不会再次调用
    MinerU；完整 Markdown 没有稳定页码，因此不添加 page 元数据。
    """
    if normalize_document_parser_mode(filename, parser_mode) != PDF_PARSER_MINERU_CLOUD:
        return []
    markdown = mineru_client.read_saved_markdown(file_path)
    if not markdown:
        return []
    return [
        Document(
            page_content=markdown,
            metadata={"source": filename, "parser_mode": PDF_PARSER_MINERU_CLOUD},
        )
    ]


def _load_pdf(
    file_path: str,
    filename: str,
    parser_mode: str | None,
    *,
    reuse_mineru_result: bool = False,
) -> list[Document]:
    """按用户显式选择分发 PDF 解析器，绝不悄悄切换到其他解析服务。"""
    mode = normalize_pdf_parser_mode(parser_mode)
    if mode == PDF_PARSER_NATIVE_TEXT:
        return _load_pdf_native_text(file_path, filename)
    if mode == PDF_PARSER_MINERU_CLOUD:
        return _load_pdf_mineru_cloud(
            file_path, filename, reuse_saved_result=reuse_mineru_result
        )
    if mode == PDF_PARSER_PADDLE_OCR:
        return _load_pdf_paddle_ocr(file_path, filename)
    raise DocumentLoadError("未知的 PDF 解析方式")


def read_text_document(file_path: str, filename: str) -> str:
    """依次尝试常见中文编码读取文本，无法识别时明确报错而非返回空内容。"""
    last_error: UnicodeError | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with open(file_path, encoding=encoding) as source:
                return source.read().replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeError as exc:
            last_error = exc
    raise DocumentLoadError(f"{filename} 文本编码无法识别，请转换为 UTF-8 或 GB18030") from last_error


def _docx_heading_level(paragraph) -> int | None:
    """从 Word 段落样式推断标题层级；文档标题固定保留为一级标题。"""
    style = getattr(paragraph, "style", None)
    style_name = str(getattr(style, "name", "") or "").strip()
    style_id = str(getattr(style, "style_id", "") or "").strip()
    if style_name.lower() == "title" or style_id.lower() == "title":
        return 1
    match = re.search(r"(?:heading|标题)\s*([1-6])", f"{style_name} {style_id}", re.IGNORECASE)
    if not match:
        return None
    # 一级标题留给 Word 文档标题，普通章节从二级开始，层级关系才不会颠倒。
    return min(6, int(match.group(1)) + 1)


def _docx_style_marker(paragraph) -> str:
    """返回用于识别本地化 Word 样式的规范化名称。"""
    style = getattr(paragraph, "style", None)
    style_name = str(getattr(style, "name", "") or "")
    style_id = str(getattr(style, "style_id", "") or "")
    return f"{style_name} {style_id}".lower()


def _docx_paragraph_is_emphasized(paragraph) -> bool:
    """识别手工标题常见的整段粗体或显式大字号信号。"""
    runs = [run for run in getattr(paragraph, "runs", ()) if str(run.text).strip()]
    if not runs:
        return False
    visible_chars = sum(len(str(run.text).strip()) for run in runs)
    bold_chars = sum(len(str(run.text).strip()) for run in runs if run.bold is True)
    mostly_bold = visible_chars > 0 and bold_chars / visible_chars >= 0.8
    explicit_sizes = [
        float(run.font.size.pt)
        for run in runs
        if getattr(getattr(run, "font", None), "size", None) is not None
    ]
    return mostly_bold or bool(explicit_sizes and max(explicit_sizes) >= 14)


def _docx_inferred_heading_level(
    paragraph, text: str, *, has_chapter_heading: bool
) -> int | None:
    """在整篇缺少标题样式时，保守推断短编号段落的章节层级。"""
    if len(text) > 80 or re.search(r"[。！？；;]$", text):
        return None
    style_marker = _docx_style_marker(paragraph)
    if "list " in style_marker or "项目符号" in style_marker or "编号" in style_marker:
        return None
    emphasized = _docx_paragraph_is_emphasized(paragraph)
    if CHAPTER_HEADING.match(text) or APPENDIX_HEADING.match(text):
        return 2
    if APPENDIX_CHILD_HEADING.match(text):
        return 3
    standard = STANDARD_HEADING.match(text)
    if standard:
        return min(6, len(standard.group(1).split(".")) + 1)
    if STANDARD_TOP_HEADING.match(text) and emphasized:
        return 2
    if CHINESE_MAJOR_HEADING.match(text) and emphasized:
        return 3 if has_chapter_heading else 2
    if CHINESE_SUB_HEADING.match(text) and emphasized:
        return 4 if has_chapter_heading else 3
    return None


def _docx_table_text(table) -> tuple[str, str]:
    """同时生成表格的纯文本版本和 HTML 版本，分别用于固定与结构化路线。"""
    plain_rows: list[str] = []
    html_rows: list[str] = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if not any(cells):
            continue
        plain_rows.append(" | ".join(cells))
        html_rows.append(
            "<tr>" + "".join(f"<td>{html.escape(cell)}</td>" for cell in cells) + "</tr>"
        )
    return "\n".join(plain_rows), f"<table>{''.join(html_rows)}</table>" if html_rows else ""


def _load_docx_local(file_path: str, filename: str) -> DocumentRoutes:
    """保留 Word 正文顺序、标题、列表和表格行，并输出两条检索路线。"""
    try:
        import docx
        from docx.table import Table

        word_document = docx.Document(file_path)
        document_blocks = list(word_document.iter_inner_content())
        has_styled_heading = any(
            not isinstance(block, Table) and _docx_heading_level(block) is not None
            for block in document_blocks
        )
        has_chapter_heading = any(
            not isinstance(block, Table)
            and CHAPTER_HEADING.match(" ".join(block.text.split()))
            for block in document_blocks
        )
        plain_blocks: list[str] = []
        structured_blocks: list[str] = []
        for block in document_blocks:
            if isinstance(block, Table):
                plain_table, structured_table = _docx_table_text(block)
                if plain_table:
                    plain_blocks.append(plain_table)
                    structured_blocks.append(structured_table)
                continue

            text = " ".join(block.text.split())
            if not text:
                continue
            plain_blocks.append(text)
            heading_level = _docx_heading_level(block)
            if heading_level is None and not has_styled_heading:
                heading_level = _docx_inferred_heading_level(
                    block, text, has_chapter_heading=has_chapter_heading
                )
            if heading_level:
                structured_blocks.append(f"{'#' * heading_level} {text}")
                continue
            style_marker = _docx_style_marker(block)
            if "list bullet" in style_marker or "项目符号" in style_marker:
                structured_blocks.append(f"- {text}")
            elif "list number" in style_marker or "编号" in style_marker:
                structured_blocks.append(f"1. {text}")
            else:
                structured_blocks.append(text)
    except Exception as exc:
        raise DocumentLoadError(f"DOCX 文本提取失败：{exc}") from exc

    plain_text = "\n\n".join(plain_blocks).strip()
    structured_text = "\n\n".join(structured_blocks).strip()
    metadata = {"source": filename, "parser_mode": PDF_PARSER_NATIVE_TEXT}
    return DocumentRoutes(
        fixed=[Document(page_content=plain_text, metadata=metadata.copy())] if plain_text else [],
        structured=[
            Document(
                page_content=structured_text,
                metadata={**metadata, "structured_profile": "native"},
            )
        ]
        if structured_text
        else [],
    )


def docx_preview_html(file_path: str) -> str:
    """按正文顺序生成安全转义的 DOCX 预览 HTML，包含标题和表格。"""
    try:
        import docx
        from docx.table import Table

        blocks: list[str] = []
        for block in docx.Document(file_path).iter_inner_content():
            if isinstance(block, Table):
                rows = []
                for row in block.rows:
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    if any(cells):
                        rows.append(
                            "<tr>"
                            + "".join(f"<td>{html.escape(cell)}</td>" for cell in cells)
                            + "</tr>"
                        )
                if rows:
                    blocks.append(f"<table>{''.join(rows)}</table>")
                continue
            text = " ".join(block.text.split())
            if not text:
                continue
            level = _docx_heading_level(block)
            tag = f"h{min(level or 0, 4)}" if level else "p"
            blocks.append(f"<{tag}>{html.escape(text)}</{tag}>")
        return "".join(blocks)
    except Exception as exc:
        raise DocumentLoadError(f"DOCX 预览生成失败：{exc}") from exc


def _load_docx_mineru_cloud(
    file_path: str,
    filename: str,
    *,
    reuse_saved_result: bool = False,
) -> DocumentRoutes:
    """使用 MinerU 解析 DOCX，并构造按页固定路线和完整 Markdown 结构化路线。"""
    try:
        pages = mineru_client.read_saved_page_texts(file_path) if reuse_saved_result else []
        if not pages:
            pages = mineru_client.parse_document(
                file_path, result_dir=mineru_client.result_dir_for_document(file_path)
            )
    except mineru_client.MinerUError as exc:
        raise DocumentLoadError(str(exc)) from exc
    fixed = [
        Document(
            page_content=page_text,
            metadata={
                "source": filename,
                "page": page_number,
                "parser_mode": PDF_PARSER_MINERU_CLOUD,
            },
        )
        for page_number, page_text in pages
    ]
    structured = load_mineru_markdown_document(
        file_path, filename, PDF_PARSER_MINERU_CLOUD
    )
    for document in structured:
        document.metadata["structured_profile"] = "mineru"
    return DocumentRoutes(fixed=fixed, structured=structured)


def load_document_routes(
    file_path: str,
    filename: str,
    parser_mode: str | None = DEFAULT_PDF_PARSER_MODE,
    *,
    reuse_mineru_result: bool = False,
) -> DocumentRoutes:
    """将一个支持的源文件加载为固定路线和可选的结构化路线。

    每个文件只选择一次解析器；同一份结果同时服务于稳定的固定切片召回，以及
    Markdown、MinerU、DOCX 的结构化检索，无需重复发起外部解析请求。
    """
    extension = os.path.splitext(filename)[1].lower()
    normalized_mode = normalize_document_parser_mode(filename, parser_mode)
    try:
        if extension == ".pdf":
            fixed = _load_pdf(
                file_path,
                filename,
                normalized_mode,
                reuse_mineru_result=reuse_mineru_result,
            )
            structured = load_mineru_markdown_document(file_path, filename, normalized_mode)
            for document in structured:
                document.metadata["structured_profile"] = "mineru"
            return DocumentRoutes(fixed=fixed, structured=structured)
        if extension == ".docx":
            if normalized_mode == PDF_PARSER_MINERU_CLOUD:
                return _load_docx_mineru_cloud(
                    file_path, filename, reuse_saved_result=reuse_mineru_result
                )
            return _load_docx_local(file_path, filename)
        if extension in (".md", ".txt"):
            text = read_text_document(file_path, filename)
            metadata = {"source": filename, "parser_mode": PDF_PARSER_NATIVE_TEXT}
            fixed = [Document(page_content=text, metadata=metadata.copy())] if text.strip() else []
            structured = []
            if extension == ".md" and text.strip():
                structured = [
                    Document(
                        page_content=text,
                        metadata={**metadata, "structured_profile": "native"},
                    )
                ]
            return DocumentRoutes(fixed=fixed, structured=structured)
        if extension in IMAGE_EXTENSIONS:
            if normalized_mode == IMAGE_PARSER_VISION_MODEL:
                from services.documents.visual import load_standalone_image_document

                document = load_standalone_image_document(
                    file_path, filename, normalized_mode
                )
                return DocumentRoutes(fixed=[document] if document is not None else [], structured=[])
            ocr = get_ocr_engine()
            if not ocr:
                raise DocumentLoadError("PaddleOCR 不可用，请检查本地 OCR 依赖")
            with Image.open(file_path) as image:
                text = ocr_image(image, ocr)
            return DocumentRoutes(
                fixed=[
                    Document(
                        page_content=text,
                        metadata={"source": filename, "parser_mode": PDF_PARSER_PADDLE_OCR},
                    )
                ] if text.strip() else [],
                structured=[],
            )
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"{filename} 解析失败：{exc}") from exc
    raise DocumentLoadError(f"暂不支持解析 {extension or '无扩展名'} 文件")
